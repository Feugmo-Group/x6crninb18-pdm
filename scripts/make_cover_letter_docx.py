# SPDX-FileCopyrightText: Copyright (c) 2023 - 2026 NVIDIA CORPORATION & AFFILIATES.
# SPDX-FileCopyrightText: All rights reserved.
# SPDX-License-Identifier: Apache-2.0

"""cover_letter.docx -- the cover letter as Word, laid out as a letter.

Nature Portfolio takes the cover letter as an editable file, so a .docx is
wanted alongside the PDF.  The LaTeX source stays the single source of truth:
this derives the Word version from it rather than keeping a second copy of the
letter that would drift the first time a sentence changed.

An earlier version piped the text through pandoc, which produced correct
content in no particular shape: the sender and recipient blocks collapsed onto
single lines because their line breaks are LaTeX \\\\, there was no date, and
the bold lead-in of each contribution was the only structure left.  This builds
the document directly instead, so the result reads as a letter.

Writes paper/cover_letter.docx.

Run:  python scripts/make_cover_letter_docx.py
"""

from __future__ import annotations

import datetime as dt
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from htw_pdm.paths import PAPER as P

SRC = P / "cover_letter.tex"
DOCX = P / "cover_letter.docx"
tex = SRC.read_text(encoding="utf-8")


def braced(cmd: str, s: str) -> str:
    i = s.index(f"\\{cmd}{{") + len(cmd) + 2
    d = 1
    for j in range(i, len(s)):
        if s[j] == "{":
            d += 1
        elif s[j] == "}":
            d -= 1
            if d == 0:
                return s[i:j]
    raise ValueError(cmd)


def detex(t: str) -> str:
    """LaTeX to plain text, keeping ** and * as inline bold/italic marks."""
    t = re.sub(r"(?<!\\)%.*", "", t)
    t = t.replace("$^\\circ$C", "\u00b0C").replace("O$_2$", "O\u2082")
    t = t.replace("$b_3$", "b\u2083").replace("$\\chi^2$", "\u03c7\u00b2")
    t = re.sub(r"\\,", "\u202f", t).replace("~", " ")
    for _ in range(6):
        before = t
        t = re.sub(r"\\textbf\{([^{}]*)\}", r"**\1**", t)
        t = re.sub(r"\\(?:textit|emph)\{([^{}]*)\}", r"*\1*", t)
        t = re.sub(r"\\texttt\{([^{}]*)\}", r"\1", t)
        if t == before:
            break
    t = t.replace("``", "\u201c").replace("''", "\u201d")
    t = t.replace("\\%", "%").replace("\\&", "&").replace("\\_", "_")
    t = t.replace("--", "\u2013")
    t = re.sub(r"\\[a-zA-Z]+", "", t)
    return re.sub(r"\s+", " ", t).strip()


def emit(par, text: str) -> None:
    """Write text into a paragraph, honouring ** bold ** and * italic * marks."""
    for piece in re.split(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not piece:
            continue
        if piece.startswith("***") and piece.endswith("***"):
            r = par.add_run(piece[3:-3])
            r.bold = r.italic = True
        elif piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


sender = [detex(x) for x in braced("address", tex).split("\\\\") if detex(x)]
recipient = [detex(x) for x in
             tex[tex.index("\\begin{letter}{") + 15: tex.index("}\n\n\\opening")].split("\\\\")
             if detex(x)]
opening = detex(braced("opening", tex))
closing = detex(braced("closing", tex))
signature = detex(braced("signature", tex))

body = tex[tex.index("\\opening"):tex.index("\\closing")]
body = body[body.index("\n"):]
body = body.replace("\\begin{itemize}", "").replace("\\end{itemize}", "")
body = re.sub(r"\\item\s+", "\n\n\u0007", body)          # \a marks a bullet
chunks = [c for c in (x.strip() for x in re.split(r"\n\s*\n", body)) if c]

doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.06
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(0.9)


def block(lines, *, bold_first=False, grey=False, after=Pt(1)):
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = after
        emit(p, ln)
        for r in p.runs:
            if bold_first and i == 0:
                r.bold = True
            if grey:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


block(sender, bold_first=True, grey=True)
d = doc.add_paragraph()
d.paragraph_format.space_before, d.paragraph_format.space_after = Pt(16), Pt(14)
d.add_run(dt.date.today().strftime("%d %B %Y")).font.color.rgb = RGBColor(0x44, 0x44, 0x44)
block(recipient)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.add_run(opening)

for c in chunks:
    if c.startswith("\u0007"):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Inches(0.28)
        emit(p, detex(c[1:]))
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        emit(p, detex(c))

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(30)      # room for a signature
p.add_run(closing)
doc.add_paragraph().add_run(signature).bold = True

doc.save(DOCX)
print(f"Saved: {DOCX}")
print(f"  {len(sender)}-line sender block, dated, {len(recipient)}-line recipient block, "
      f"{sum(1 for c in chunks if c.startswith(chr(7)))} bulleted contributions")
